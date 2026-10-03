import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

print("=" * 70)
print("生成最终Excel报告")
print("=" * 70)

# 读取所有数据
print("\n读取数据...")
with open('tube_entities_v9.json', 'r', encoding='utf-8') as f:
    step1_tubes = json.load(f)

with open('step2_contacts_v9.json', 'r', encoding='utf-8') as f:
    step2_contacts = json.load(f)

with open('step3_positions_v9.json', 'r', encoding='utf-8') as f:
    step3_positions = json.load(f)

with open('step4_split_points_v9.json', 'r', encoding='utf-8') as f:
    step4_splits = json.load(f)

with open('step5_merge_decisions_v9.json', 'r', encoding='utf-8') as f:
    step5_decisions = json.load(f)

with open('step6_topology_v9.json', 'r', encoding='utf-8') as f:
    step6_topology = json.load(f)

with open('step7_attachment_v9.json', 'r', encoding='utf-8') as f:
    step7_attachments = json.load(f)

# 创建 Excel 工作簿
wb = openpyxl.Workbook()

# 样式定义
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF", size=11)
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)
center_align = Alignment(horizontal="center", vertical="center")
left_align = Alignment(horizontal="left", vertical="center")

def style_header(ws, headers):
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

# ==================== 第1步：管状实体 ====================
print("创建工作表1: 管状实体...")
ws1 = wb.active
ws1.title = "1-管状实体"

headers1 = ["序号", "实体名称", "总长度(mm)", "圆柱长度(mm)", "弯曲长度(mm)", 
            "半径(mm)", "起点X", "起点Y", "起点Z", "终点X", "终点Y", "终点Z"]
ws1.append(headers1)
style_header(ws1, headers1)

for i, tube in enumerate(step1_tubes, 1):
    row_data = [
        i,
        tube['name'],
        round(tube['length'], 1),
        round(tube['cyl_length'], 1),
        round(tube['bspline_length'], 1),
        round(tube['radius'], 1),
        round(tube['p0'][0], 1),
        round(tube['p0'][1], 1),
        round(tube['p0'][2], 1),
        round(tube['p1'][0], 1),
        round(tube['p1'][1], 1),
        round(tube['p1'][2], 1)
    ]
    ws1.append(row_data)

ws1.column_dimensions['A'].width = 6
ws1.column_dimensions['B'].width = 60
for col in 'CDEFGHIJKL':
    ws1.column_dimensions[col].width = 12

# ==================== 第2步：接触关系 ====================
print("创建工作表2: 接触关系...")
ws2 = wb.create_sheet("2-接触关系")

headers2 = ["序号", "实体名称", "实体中心X", "实体中心Y", "实体中心Z",
            "接触管状实体", "接触端点", "接触点X", "接触点Y", "接触点Z", "距离(mm)"]
ws2.append(headers2)
style_header(ws2, headers2)

for i, contact in enumerate(step2_contacts, 1):
    if not contact['contacts']:
        row_data = [
            i, contact['entity'],
            round(contact['center'][0], 1),
            round(contact['center'][1], 1),
            round(contact['center'][2], 1),
            "无接触", "", "", "", "", ""
        ]
        ws2.append(row_data)
    else:
        first = True
        for c in contact['contacts']:
            if first:
                row_data = [
                    i, contact['entity'],
                    round(contact['center'][0], 1),
                    round(contact['center'][1], 1),
                    round(contact['center'][2], 1),
                    c['tube'], c['end'],
                    round(c['contact_point'][0], 1),
                    round(c['contact_point'][1], 1),
                    round(c['contact_point'][2], 1),
                    round(c['distance'], 1)
                ]
                ws2.append(row_data)
                first = False
            else:
                row_data = [
                    "", "", "", "", "",
                    c['tube'], c['end'],
                    round(c['contact_point'][0], 1),
                    round(c['contact_point'][1], 1),
                    round(c['contact_point'][2], 1),
                    round(c['distance'], 1)
                ]
                ws2.append(row_data)

ws2.column_dimensions['A'].width = 6
ws2.column_dimensions['B'].width = 60
for col in 'CDEFGHIJK':
    ws2.column_dimensions[col].width = 12

# ==================== 第3步：定位坐标 ====================
print("创建工作表3: 定位坐标...")
ws3 = wb.create_sheet("3-定位坐标")

headers3 = ["序号", "实体名称", "定位坐标X", "定位坐标Y", "定位坐标Z",
            "接触数", "推理原因"]
ws3.append(headers3)
style_header(ws3, headers3)

for i, pos in enumerate(step3_positions, 1):
    row_data = [
        i, pos['name'],
        round(pos['position'][0], 1),
        round(pos['position'][1], 1),
        round(pos['position'][2], 1),
        pos['contact_count'],
        pos['reason']
    ]
    ws3.append(row_data)

ws3.column_dimensions['A'].width = 6
ws3.column_dimensions['B'].width = 60
for col in 'CDE':
    ws3.column_dimensions[col].width = 12
ws3.column_dimensions['F'].width = 10
ws3.column_dimensions['G'].width = 50

# ==================== 第4步：分段点 ====================
print("创建工作表4: 分段点...")
ws4 = wb.create_sheet("4-分段点")

headers4 = ["分段点ID", "类型", "关联实体", "位置X", "位置Y", "位置Z", "备注"]
ws4.append(headers4)
style_header(ws4, headers4)

for sp in step4_splits:
    note = ""
    if 'nearby_ends' in sp:
        note = f"交汇{len(sp['nearby_ends'])}个端点"
    row_data = [
        sp['point_id'],
        sp['split_reason'],
        sp['entity'],
        round(sp['position'][0], 1),
        round(sp['position'][1], 1),
        round(sp['position'][2], 1),
        note
    ]
    ws4.append(row_data)

ws4.column_dimensions['A'].width = 12
ws4.column_dimensions['B'].width = 12
ws4.column_dimensions['C'].width = 60
for col in 'DEF':
    ws4.column_dimensions[col].width = 12
ws4.column_dimensions['G'].width = 20

# ==================== 第5步：合并决策 ====================
print("创建工作表5: 合并决策...")
ws5 = wb.create_sheet("5-合并决策")

headers5 = ["序号", "管状实体A", "端点A", "管状实体B", "端点B", 
            "端点距离(mm)", "是否有分段点", "分段实体", "决策", "原因"]
ws5.append(headers5)
style_header(ws5, headers5)

for i, decision in enumerate(step5_decisions, 1):
    has_split = "是" if decision['has_split_point'] else "否"
    split_entity = decision['split_entity'] if decision['split_entity'] else ""
    row_data = [
        i,
        decision['tube_a'],
        decision['end_a'],
        decision['tube_b'],
        decision['end_b'],
        round(decision['end_distance'], 1),
        has_split,
        split_entity,
        decision['decision'],
        decision['reason']
    ]
    ws5.append(row_data)

ws5.column_dimensions['A'].width = 6
ws5.column_dimensions['B'].width = 60
ws5.column_dimensions['C'].width = 8
ws5.column_dimensions['D'].width = 60
ws5.column_dimensions['E'].width = 8
for col in 'FGHIJ':
    ws5.column_dimensions[col].width = 15

# ==================== 第6步：拓扑关系 ====================
print("创建工作表6: 拓扑关系...")
ws6 = wb.create_sheet("6-拓扑关系")

# 节点表
ws6.append(["=== 节点列表 ==="])
ws6.cell(1, 1).font = Font(bold=True, size=12)

headers6a = ["节点ID", "位置X", "位置Y", "位置Z", "成员数", "连接线束段"]
ws6.append(headers6a)
style_header(ws6, headers6a)

for node in step6_topology['nodes']:
    connected = ", ".join(node['connected_tubes'])
    row_data = [
        node['node_id'],
        round(node['position'][0], 1),
        round(node['position'][1], 1),
        round(node['position'][2], 1),
        node['member_count'],
        connected
    ]
    ws6.append(row_data)

# 线束段表
start_row = len(step6_topology['nodes']) + 3
ws6.cell(start_row, 1, "=== 线束段列表 ===")
ws6.cell(start_row, 1).font = Font(bold=True, size=12)

headers6b = ["管状实体", "总长度(mm)", "圆柱长度(mm)", "弯曲长度(mm)", "半径(mm)",
             "起点X", "起点Y", "起点Z", "终点X", "终点Y", "终点Z", "起始节点", "终止节点"]
for col, header in enumerate(headers6b, 1):
    cell = ws6.cell(start_row + 1, col, header)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = center_align
    cell.border = thin_border

for i, seg in enumerate(step6_topology['segments'], 1):
    row = start_row + 1 + i
    ws6.cell(row, 1, seg['tube'])
    ws6.cell(row, 2, round(seg['length'], 1))
    ws6.cell(row, 3, round(seg['cyl_length'], 1))
    ws6.cell(row, 4, round(seg['bspline_length'], 1))
    ws6.cell(row, 5, round(seg['radius'], 1))
    ws6.cell(row, 6, round(seg['p0'][0], 1))
    ws6.cell(row, 7, round(seg['p0'][1], 1))
    ws6.cell(row, 8, round(seg['p0'][2], 1))
    ws6.cell(row, 9, round(seg['p1'][0], 1))
    ws6.cell(row, 10, round(seg['p1'][1], 1))
    ws6.cell(row, 11, round(seg['p1'][2], 1))
    ws6.cell(row, 12, seg['node_start'])
    ws6.cell(row, 13, seg['node_end'])
    
    for col in range(1, 14):
        ws6.cell(row, col).border = thin_border
        ws6.cell(row, col).alignment = center_align

ws6.column_dimensions['A'].width = 60
for col in 'BCDEFGHIJKLM':
    ws6.column_dimensions[col].width = 12

# ==================== 第7步：实体定位 ====================
print("创建工作表7: 实体定位...")
ws7 = wb.create_sheet("7-实体定位")

headers7 = ["序号", "实体名称", "类型", "定位坐标X", "定位坐标Y", "定位坐标Z",
            "挂在线束段", "参数t", "距起点距离(mm)", "距线束段距离(mm)"]
ws7.append(headers7)
style_header(ws7, headers7)

for i, att in enumerate(step7_attachments, 1):
    row_data = [
        i,
        att['name'],
        att['type'],
        round(att['position'][0], 1),
        round(att['position'][1], 1),
        round(att['position'][2], 1),
        att['attached_segment'],
        round(att['parameter_t'], 3),
        round(att['distance_from_start'], 1),
        round(att['distance_to_segment'], 1)
    ]
    ws7.append(row_data)

ws7.column_dimensions['A'].width = 6
ws7.column_dimensions['B'].width = 60
ws7.column_dimensions['C'].width = 12
for col in 'DEFGHIJ':
    ws7.column_dimensions[col].width = 15

# 保存文件
output_file = "线束拓扑分析_7步过程数据_v9.xlsx"
wb.save(output_file)

print(f"\n已保存: {output_file}")
print(f"  工作表1: 管状实体 ({len(step1_tubes)} 个)")
print(f"  工作表2: 接触关系 ({len(step2_contacts)} 个实体)")
print(f"  工作表3: 定位坐标 ({len(step3_positions)} 个)")
print(f"  工作表4: 分段点 ({len(step4_splits)} 个)")
print(f"  工作表5: 合并决策 ({len(step5_decisions)} 对)")
print(f"  工作表6: 拓扑关系 ({len(step6_topology['nodes'])} 节点, {len(step6_topology['segments'])} 线束段)")
print(f"  工作表7: 实体定位 ({len(step7_attachments)} 个)")
