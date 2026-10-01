#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""新版全流程合成测试: build_topology_from_relations -> make_report -> 校验 xlsx."""
import sys, math
from unittest.mock import MagicMock
for name in ['OCP.STEPCAFControl', 'OCP.TDocStd', 'OCP.TCollection', 'OCP.XCAFDoc',
             'OCP.collections', 'OCP.TDF', 'OCP.TDataStd', 'OCP.TopExp', 'OCP.TopAbs',
             'OCP.TopoDS', 'OCP.BRepAdaptor', 'OCP.BRepGProp', 'OCP.GProp', 'OCP.Bnd',
             'OCP.BRepBndLib', 'OCP.TopLoc', 'OCP.GeomAbs', 'OCP.gp', 'OCP.BRepExtrema', 'OCP.BRepBuilderAPI']:
    sys.modules[name] = MagicMock()
import numpy as np

SRC = '/home/hatch/workspace/harness_3d/harness_topology.py'
ns = {'__name__': 't'}
exec(compile(open(SRC).read(), SRC, 'exec'), ns)
build = ns['build_topology_from_relations']
_d3 = ns['_d3']

sys.path.insert(0, '/home/hatch/workspace/harness_3d')
import report_xlsx

def mkbranch(key, pts, occ=None, dia=10.0):
    pts = [[float(x) for x in p] for p in pts]
    L = sum(_d3(pts[i], pts[i+1]) for i in range(len(pts)-1))
    return {"key": key, "proto": key, "occ": occ or key, "n_seg": 1,
            "length": L, "vol": None, "dia": dia,
            "p0": pts[0], "p1": pts[-1], "pts": pts, "src": "reversed", "solids": []}

# 场景: 主干 MAIN(0..100) + TAP1 在 s=30 搭接 + TAP2 在 s=33 搭接(合并) + 扎带站位 s=70 + 连接器在 MAIN 末端
M = mkbranch("MAIN", [[0,0,0],[100,0,0]], occ="TUBE\\MAIN", dia=12.0)
T1 = mkbranch("TAP1", [[30,0,0],[30,40,0]], occ="TUBE\\TAP1", dia=8.0)
T2 = mkbranch("TAP2", [[33,0,0],[33,-40,0]], occ="TUBE\\TAP2", dia=8.0)
rels = {
    "taps": [{"main": 0, "s": 30.0, "xyz": (30,0,0), "tap": 1, "tap_end": 0, "dev": 0.4},
             {"main": 0, "s": 33.0, "xyz": (33,0,0), "tap": 2, "tap_end": 0, "dev": 0.4}],
    "end_ends": [],
    "terminals": [{"branch": 0, "end": 1, "tag": "CONNECTOR\\AXA\\test", "kind": "connector", "dist": 0.2}],
    "tie_stations": [{"branch": 0, "s": 70.0, "xyz": (70,0,0), "tag": "TIE\\Z01", "kind": "tie", "dev": 0.2}],
}
segs, nodes, ents, runs, bi = build([M, T1, T2], rels, 3.0, [], lambda *a: None)

branches = []
for i, b in enumerate([M, T1, T2]):
    branches.append({"key": b["key"], "proto": b["proto"], "occ": b["occ"], "n_seg": 1,
                     "length": round(b["length"],1), "dia": b["dia"], "src": b["src"],
                     "p0": b["p0"], "p1": b["p1"],
                     "polyline": b["pts"],
                     "segs": bi[i]["segs"], "node0": bi[i]["node0"], "node1": bi[i]["node1"]})
d = {"file": "SYNTH.stp", "node_tol": 3.0, "branches": branches,
     "segments": segs, "nodes": nodes, "entities": ents, "runs": runs,
     "connectors": [{"occ": "CONNECTOR\\AXA\\test", "proto": "AXA",
                     "center": [100,0,0], "bbox": [20,20,20]},
                    {"occ": "TIE\\Z01", "proto": "Z01",
                     "center": [70,5,0], "bbox": [8,8,8]}]}
report_xlsx.make_report(d, "/tmp/synth_report.xlsx")
print("report saved")

from openpyxl import load_workbook
wb = load_workbook("/tmp/synth_report.xlsx")
print("sheets:", wb.sheetnames)
FAIL = 0
def check(name, cond, extra=""):
    global FAIL
    print(("  PASS " if cond else "  FAIL ") + name, extra if not cond else "")
    if not cond: FAIL += 1

ws = wb["线段表"]
rows = list(ws.iter_rows(min_row=2, values_only=True))
check("线段数=5(主干3段+2分支)", len(rows) == 5, f"n={len(rows)}")
for r in rows: print("   ", r[0], r[1], "->", r[2], "L=", r[3], "tube=", r[6], "run=", r[7])

ws = wb["节点表"]
nrows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
for k, r in nrows.items(): print("   ", r[0], r[2], "3d=", r[1], "segs=", r[7])
check("BN01 存在且为分支点", "BN01" in nrows and nrows["BN01"][3] == "分支点")
check("BN01 度数=4(TAP1+TAP2合并)", nrows.get("BN01", [0]*10)[7] == 4, f"{nrows.get('BN01')}")
check("CON01 3D名正确", nrows.get("CON01", [0, ""])[1] == "CONNECTOR\\AXA\\test")
check("CLP01 3D名正确", nrows.get("CLP01", [0, ""])[1] == "TIE\\Z01")
check("CLP01 类型为固定卡扣", nrows.get("CLP01", [0,0,0,""])[3] == "固定卡扣")

ws = wb["编码对照表"]
erows = list(ws.iter_rows(min_row=2, values_only=True))
ec = {r[0]: r for r in erows}
check("对照表: CON01->CONNECTOR\\AXA\\test",
      ec.get("CON01", ["", "", ""])[2] == "CONNECTOR\\AXA\\test")
check("对照表: SEG01 STEP标签含TUBE", "TUBE" in str(ec.get("SEG01", ["", "", "", ""])[3]))
print("   对照表条目:", len(erows))

ws = wb["连续走线"]
rrows = list(ws.iter_rows(min_row=2, values_only=True))
for r in rrows: print("   ", r[0], r[1], "L=", r[2], r[3], "->", r[4])
# 走线应在 BN01 处截断, 但穿过度数2的 CLP01(卡扣点保留在线上): 共4条
check("走线数=4(CLP01不断开)", len(rrows) == 4, f"n={len(rrows)}")
r2 = [r for r in rrows if r[0] == "走线2"][0]
check("走线2 穿过CLP01", r2[1] == "SEG02 → SEG03" and r2[3] == "BN01" and r2[4] == "CON01",
      f"{r2[1]}")

ws = wb["连接器卡扣清单"]
crows = list(ws.iter_rows(min_row=2, values_only=True))
cc = {r[1]: r[0] for r in crows}
check("连接器清单带编码", cc.get("CONNECTOR\\AXA\\test") == "CON01"
      and cc.get("TIE\\Z01") == "CLP01", f"{cc}")

print("FAIL:", FAIL)
sys.exit(1 if FAIL else 0)
